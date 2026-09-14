import csv
import hashlib
import io
import json
import zipfile
from decimal import Decimal, InvalidOperation
from typing import Any
from uuid import uuid5

import openpyxl
from django.core.exceptions import ValidationError
from django.db import transaction

from app.accounts.policies import require_operator
from app.catalog.models import SKU
from app.common.business import BusinessError, record_event
from app.common.services import IdempotencyConflict
from app.inventory.models import StockLot
from app.orders.models import SalesOrder
from app.orders.services import create_order
from app.shops.models import SalesChannel

from .models import ImportJob, ImportRow

HEADERS = ["渠道订单号", "渠道代码", "商品编码", "客户称呼", "数量", "成交单价", "实物组ID"]


def preview_import(*, actor, upload, column_mapping=None):
    require_operator(actor)
    if not 0 < upload.size <= 10 * 1024 * 1024:
        raise BusinessError("文件需在 10 MB 以内。")
    raw = upload.read()
    mapping = {header: (column_mapping or {}).get(header) or header for header in HEADERS}
    fingerprint = hashlib.sha256(
        raw + json.dumps(mapping, sort_keys=True, ensure_ascii=False).encode()
    ).hexdigest()
    old = ImportJob.objects.filter(file_hash=fingerprint).first()
    if old:
        return old
    rows: Any
    if upload.name.lower().endswith(".csv"):
        try:
            content = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            content = raw.decode("gb18030")
        rows = csv.reader(io.StringIO(content))
    elif upload.name.lower().endswith(".xlsx"):
        try:
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                if sum(info.file_size for info in archive.infolist()) > 50 * 1024 * 1024:
                    raise BusinessError("Excel 解压后过大，请拆分文件。")
            workbook = openpyxl.load_workbook(
                io.BytesIO(raw), read_only=True, data_only=False, keep_links=False
            )
            sheet = workbook.active
            if sheet.max_column > 30 or sheet.max_row > 10001:
                workbook.close()
                raise BusinessError("每次最多 10000 行、30 列，请拆分文件。")
            rows = list(sheet.iter_rows(values_only=True))
            workbook.close()
        except (zipfile.BadZipFile, KeyError, ValueError) as exc:
            raise BusinessError("Excel 文件损坏或不受支持。") from exc
        rows = iter(rows)
    else:
        raise BusinessError("请选择 CSV 或 XLSX 文件。")
    headers = [str(value or "").strip() for value in next(rows, [])]
    reverse_mapping = {value: key for key, value in mapping.items()}
    if len(reverse_mapping) != len(mapping):
        raise BusinessError("不同字段不能对应同一列，请检查列名对应关系。")
    headers = [reverse_mapping.get(header, header) for header in headers]
    if not all(header in headers for header in HEADERS[:-1]) or len(set(headers)) != len(headers):
        raise BusinessError("请使用标准模板表头：" + "、".join(HEADERS))
    prepared = []
    for number, values in enumerate(rows, 2):
        if number > 10001:
            raise BusinessError("每次最多导入 10000 行。")
        if not any(value is not None and str(value).strip() for value in values):
            continue
        data = dict(
            zip(
                headers,
                [str(value).strip() if value is not None else "" for value in values],
                strict=False,
            )
        )
        try:
            quantity = int(data.get("数量", ""))
            price = Decimal(data.get("成交单价", ""))
            if (
                not price.is_finite()
                or price < 0
                or price > 10**10
                or price * 100 != (price * 100).to_integral_value()
                or not 1 <= quantity <= 1000000
            ):
                raise ValueError
            external = data.get("渠道订单号", "")
            customer = data.get("客户称呼", "")
            if not 1 <= len(external) <= 100 or not 1 <= len(customer) <= 100:
                raise ValueError
            channel = SalesChannel.objects.filter(code=data.get("渠道代码"), is_active=True).first()
            sku = SKU.objects.filter(code=data.get("商品编码"), is_active=True).first()
            if not channel or not sku:
                raise BusinessError("找不到渠道代码或商品编码。")
            lot_id = data.get("实物组ID", "")
            if lot_id and not StockLot.objects.filter(pk=lot_id, sku=sku).exists():
                raise BusinessError("实物组不属于该商品。")
            if sku.requires_explicit_lot_selection and not lot_id:
                raise BusinessError("该商品货况不同，请填写实物组ID。")
            payload = dict(
                channel_id=str(channel.pk),
                sku_id=str(sku.pk),
                quantity=quantity,
                unit_price_fen=int(price * 100),
                external_order_no=external,
                customer_name=customer,
                selected_lot_id=lot_id or None,
            )
            prepared.append(ImportRow(number=number, payload=payload))
        except (ValueError, InvalidOperation, BusinessError, ValidationError) as exc:
            prepared.append(
                ImportRow(
                    number=number,
                    payload=data,
                    status="ERROR",
                    error=str(exc)[:500]
                    if isinstance(exc, BusinessError)
                    else "数量、金额或必填内容格式不正确。",
                )
            )
    if not prepared:
        raise BusinessError("文件没有可预览的数据行。")
    with transaction.atomic():
        job, created = ImportJob.objects.get_or_create(
            file_hash=fingerprint, defaults={"filename": upload.name[:200], "actor": actor}
        )
        if created:
            for row in prepared:
                row.job = job
            ImportRow.objects.bulk_create(prepared)
            record_event(actor, "import.previewed", job, rows=len(prepared))
    return job


def execute_import(job_id):
    from django.core.exceptions import ValidationError

    job = ImportJob.objects.select_related("actor").get(pk=job_id)
    if job.status not in ("QUEUED", "RUNNING"):
        return
    ImportJob.objects.filter(pk=job.pk).update(status="RUNNING")
    for row_id in job.rows.filter(status="READY").values_list("id", flat=True).iterator():
        with transaction.atomic():
            row = ImportRow.objects.select_for_update().get(pk=row_id)
            if row.status != "READY":
                continue
            try:
                payload = row.payload
                existing = SalesOrder.objects.filter(
                    channel_id=payload["channel_id"], external_order_no=payload["external_order_no"]
                ).first()
                if existing:
                    row.order = existing
                    row.status = "SKIPPED"
                else:
                    result = create_order(
                        actor=job.actor, submission_key=uuid5(job.pk, str(row.number)), **payload
                    )
                    row.order_id = result["order_id"]
                    row.status = "IMPORTED"
            except (BusinessError, IdempotencyConflict, ValidationError) as exc:
                row.status = "ERROR"
                row.error = str(exc)[:500]
            row.save()
            if row.order_id:
                channel = SalesChannel.objects.get(pk=payload["channel_id"])
                if channel.code == "XIANYU":
                    # Import and API share the same channel/order identity. Linking metadata
                    # never replays stock, fulfillment or cash events.
                    from app.integrations.models import (
                        ExternalFactApplication,
                        PlatformOrder,
                    )

                    platform_rows = PlatformOrder.objects.filter(
                        external_order_no=payload["external_order_no"],
                        order__isnull=True,
                    )
                    connection_ids = list(platform_rows.values_list("connection_id", flat=True))
                    platform_rows.update(
                        order_id=row.order_id, auto_matched=True, needs_review=True
                    )
                    ExternalFactApplication.objects.filter(
                        connection_id__in=connection_ids,
                        fact_type="ORDER",
                        external_key=payload["external_order_no"],
                    ).update(
                        result=ExternalFactApplication.Result.LINKED,
                        target_type="SalesOrder",
                        target_id=str(row.order_id),
                        reason="历史导入按渠道订单号匹配；未重放库存、履约或现金事件。",
                    )
    ImportJob.objects.filter(pk=job.pk).update(status="DONE", error="")
